"""Leitura, importação (CSV e XLSX) e gravação de `carteira/transacoes.csv` (ADR 0020, itens 3 e 4).

O arquivo é só-acréscimo: o histórico não é reescrito. Uma importação é tudo ou nada: se qualquer
linha falhar, ou o histórico resultante não fechar (venda acima da posição), nada é gravado e o
erro aponta a linha. Números aceitam o separador decimal do formato (`,` no br) e milhar com `.`
só em agrupamentos de três dígitos; `1.5` num arquivo br é erro, não 15.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from collections.abc import Iterable, Mapping, Sequence
from decimal import Decimal
from pathlib import Path

from valuation.carteira import ErroCarteira
from valuation.carteira.transacoes import Transacao, nova_transacao, validar_historico
from valuation.config import ConfigCsv

COLUNAS = ("data", "ticker", "tipo", "quantidade", "preco", "valor", "taxas", "observacao")
OBRIGATORIAS = ("data", "ticker", "tipo", "quantidade")
PADRAO_DATA_BR = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")


def _data(valor: object) -> dt.date:
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    texto = str(valor).strip()
    br = PADRAO_DATA_BR.match(texto)
    try:
        if br:
            return dt.date(int(br[3]), int(br[2]), int(br[1]))
        return dt.date.fromisoformat(texto[:10] if len(texto) > 10 else texto)
    except ValueError:
        raise ErroCarteira(f"data inválida '{texto}' (use AAAA-MM-DD ou DD/MM/AAAA)") from None


def _numero(valor: object, decimal: str) -> float | None:
    """Número da célula; vazio vira `None`. Texto segue o separador decimal do formato."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, bool):
        raise ErroCarteira(f"número inválido '{valor}'")
    if isinstance(valor, int | float | Decimal):
        return float(valor)
    texto = str(valor).strip()
    milhar = "." if decimal == "," else ","
    padrao = rf"^[+-]?(\d+|\d{{1,3}}(\{milhar}\d{{3}})+)(\{decimal}\d+)?$"
    if not re.match(padrao, texto):
        raise ErroCarteira(f"número inválido '{texto}' (decimal '{decimal}')")
    return float(texto.replace(milhar, "").replace(decimal, "."))


def _texto(valor: object) -> str:
    return "" if valor is None else str(valor).strip()


def converter_linha(celulas: Mapping[str, object], decimal: str) -> Transacao:
    """Uma linha (cabeçalho -> célula) em `Transacao`; campo inválido vira `ErroCarteira`."""
    for coluna in OBRIGATORIAS:
        if _texto(celulas.get(coluna)) == "":
            raise ErroCarteira(f"coluna '{coluna}' vazia")
    quantidade = _numero(celulas["quantidade"], decimal)
    assert quantidade is not None  # garantido por OBRIGATORIAS
    return nova_transacao(
        data=_data(celulas["data"]),
        ticker=_texto(celulas["ticker"]),
        tipo=_texto(celulas["tipo"]),
        quantidade=quantidade,
        preco=_numero(celulas.get("preco"), decimal),
        valor=_numero(celulas.get("valor"), decimal),
        taxas=_numero(celulas.get("taxas"), decimal),
        observacao=_texto(celulas.get("observacao")),
    )


def _cabecalho(nomes: Sequence[object]) -> list[str]:
    colunas = [_texto(n).lower() for n in nomes]
    desconhecidas = [c for c in colunas if c and c not in COLUNAS]
    faltando = [c for c in OBRIGATORIAS if c not in colunas]
    if desconhecidas or faltando:
        partes = []
        if faltando:
            partes.append(f"faltam as colunas {', '.join(faltando)}")
        if desconhecidas:
            partes.append(f"colunas desconhecidas: {', '.join(desconhecidas)}")
        problemas = "; ".join(partes)
        raise ErroCarteira(f"cabeçalho (linha 1): {problemas}; esperado: {', '.join(COLUNAS)}")
    return colunas


def _converter(
    linhas: Iterable[tuple[int, Sequence[object]]], colunas: list[str], decimal: str
) -> tuple[list[Transacao], list[int]]:
    transacoes: list[Transacao] = []
    numeros: list[int] = []
    erros: list[str] = []
    for numero, celulas in linhas:
        if all(_texto(c) == "" for c in celulas):
            continue
        if len(celulas) > len(colunas) and any(_texto(c) for c in celulas[len(colunas) :]):
            erros.append(f"linha {numero}: mais células que colunas no cabeçalho")
            continue
        registro = dict(zip(colunas, celulas, strict=False))
        try:
            transacoes.append(converter_linha(registro, decimal))
            numeros.append(numero)
        except ErroCarteira as e:
            erros.append(f"linha {numero}: {e}")
    if erros:
        raise ErroCarteira(erros)
    return transacoes, numeros


def _linhas_csv(caminho: Path, config: ConfigCsv) -> Iterable[tuple[int, list[str]]]:
    with caminho.open(newline="", encoding=config.encoding) as f:
        leitor = csv.reader(f, delimiter=config.separador)
        for celulas in leitor:
            yield leitor.line_num, celulas


def _linhas_xlsx(caminho: Path) -> Iterable[tuple[int, tuple[object, ...]]]:
    from openpyxl import load_workbook

    pasta = load_workbook(caminho, read_only=True, data_only=True)
    try:
        planilha = pasta.active
        if planilha is None:
            raise ErroCarteira(f"{caminho.name}: planilha sem abas")
        yield from enumerate(planilha.iter_rows(values_only=True), start=1)
    finally:
        pasta.close()


def ler_arquivo(caminho: Path, config: ConfigCsv) -> tuple[list[Transacao], list[int]]:
    """Lê CSV (`config`) ou XLSX (primeira aba, cabeçalho na linha 1).

    Devolve as transações e o número da linha de cada uma no arquivo. Qualquer linha inválida
    recusa o arquivo inteiro (`ErroCarteira` com todas as linhas e motivos).
    """
    if not caminho.exists():
        raise ErroCarteira(f"arquivo não encontrado: {caminho}")
    sufixo = caminho.suffix.lower()
    if sufixo == ".xlsx":
        linhas: Iterable[tuple[int, Sequence[object]]] = _linhas_xlsx(caminho)
    elif sufixo in (".csv", ".txt"):
        linhas = _linhas_csv(caminho, config)
    else:
        raise ErroCarteira(f"formato '{sufixo}' não suportado (use .csv ou .xlsx)")
    iterador = iter(linhas)
    primeira = next(iterador, None)
    if primeira is None:
        raise ErroCarteira(f"{caminho.name}: arquivo vazio")
    colunas = _cabecalho(primeira[1])
    return _converter(iterador, colunas, config.decimal)


def carregar(caminho: Path, config: ConfigCsv) -> tuple[list[Transacao], list[int]]:
    """Carteira gravada; arquivo ausente é carteira vazia. Valida o histórico inteiro."""
    if not caminho.exists():
        return [], []
    transacoes, numeros = ler_arquivo(caminho, config)
    validar_historico(transacoes, [f"linha {n}" for n in numeros])
    return transacoes, numeros


def _numero_texto(valor: float | None, decimal: str) -> str:
    if valor is None:
        return ""
    texto = format(Decimal(repr(round(valor, 8))), "f")
    if "." in texto:
        texto = texto.rstrip("0").rstrip(".")
    return texto.replace(".", decimal)


def _linha_csv(t: Transacao, decimal: str) -> list[str]:
    return [
        t.data.isoformat(),
        t.ticker,
        t.tipo.value,
        _numero_texto(t.quantidade, decimal),
        _numero_texto(t.preco, decimal),
        _numero_texto(t.valor if t.valor else None, decimal),
        _numero_texto(t.taxas if t.taxas else None, decimal),
        t.observacao,
    ]


def acrescentar(
    caminho: Path,
    config: ConfigCsv,
    novas: Sequence[Transacao],
    linhas_origem: Sequence[int] | None = None,
    *,
    hoje: dt.date | None = None,
    permitir_duplicadas: bool = False,
) -> int:
    """Valida carteira + `novas` e só então acrescenta ao arquivo; devolve quantas gravou.

    `linhas_origem` (linha de cada nova no arquivo importado) melhora as mensagens de erro.
    Uma nova idêntica a uma já gravada é recusada (reimportação), salvo `permitir_duplicadas`.
    """
    if not novas:
        raise ErroCarteira("nada a gravar: nenhuma transação")
    existentes, numeros = carregar(caminho, config)
    if not permitir_duplicadas:
        repetidas = [n for n in novas if n in existentes]
        if repetidas:
            raise ErroCarteira(
                f"{len(repetidas)} transação(ões) idêntica(s) a já gravadas (reimportação?), "
                f"p.ex. {repetidas[0].tipo.value} {repetidas[0].ticker} em "
                f"{repetidas[0].data.isoformat()}; use --permitir-duplicadas se for intencional"
            )
    origem = linhas_origem or range(2, len(novas) + 2)
    rotulos = [f"carteira linha {n}" for n in numeros] + [f"nova, linha {n}" for n in origem]
    validar_historico([*existentes, *novas], rotulos, hoje)

    existe = caminho.exists()
    caminho.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig escreve o BOM a cada abertura: só no arquivo novo.
    encoding = config.encoding if not existe else config.encoding.replace("-sig", "")
    precisa_quebra = False
    if existe and caminho.stat().st_size:
        with caminho.open("rb") as f:
            f.seek(-1, 2)
            precisa_quebra = f.read(1) not in (b"\n", b"\r")
    with caminho.open("a", newline="", encoding=encoding) as f:
        if precisa_quebra:
            f.write("\n")
        escritor = csv.writer(f, delimiter=config.separador, lineterminator="\n")
        if not existe or caminho.stat().st_size == 0:
            escritor.writerow(COLUNAS)
        escritor.writerows(_linha_csv(t, config.decimal) for t in novas)
    return len(novas)
